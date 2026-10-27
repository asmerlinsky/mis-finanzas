"""Redact an `.xlsx` statement into an `.xlsx` sample (rules in `tools.redaction`).

Covers every sheet (hidden ones too), sheet titles, headers and footers. Comments, hyperlinks
and document metadata are removed.
"""

import logging
import random
from collections.abc import Iterator
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.packaging.core import DocumentProperties
from openpyxl.worksheet.header_footer import HeaderFooterItem
from openpyxl.worksheet.worksheet import Worksheet

from mis_finanzas.tools.redaction import (
    MAX_ATTEMPTS,
    Redactor,
    check_paths,
    leak_error,
    text_leaks,
)

logger = logging.getLogger(__name__)

HEADER_FOOTERS = (
    "oddHeader",
    "oddFooter",
    "evenHeader",
    "evenFooter",
    "firstHeader",
    "firstFooter",
)


def _header_footer_texts(ws: Worksheet) -> Iterator[tuple[HeaderFooterItem, str, str]]:
    """`(item, "left"|"center"|"right", text)` for every non-empty header/footer part."""
    for attribute in HEADER_FOOTERS:
        item: HeaderFooterItem | None = getattr(ws, attribute)
        if item is None:
            continue
        for position in ("left", "center", "right"):
            text = getattr(item, position).text
            if text:
                yield item, position, text


def _redact_cells(ws: Worksheet, redactor: Redactor) -> None:
    for row in ws.iter_rows():
        for cell in row:
            if cell.comment is not None:
                cell.comment = None
            if cell.hyperlink is not None:
                cell.hyperlink = None
            value = cell.value
            if isinstance(value, bool) or value is None or isinstance(value, date):
                continue
            if isinstance(value, int | float):
                cell.value = redactor.number(value)
            elif isinstance(value, str):
                cell.value = redactor.text(value)


def redact_workbook(wb: Workbook, redactor: Redactor) -> None:
    """Redact in place."""
    for ws in wb.worksheets:
        _redact_cells(ws, redactor)
        ws.title = redactor.text(ws.title)
        for item, position, text in list(_header_footer_texts(ws)):
            getattr(item, position).text = redactor.text(text)
    wb.properties = DocumentProperties()


def workbook_texts(wb: Workbook) -> Iterator[str]:
    """Every value a reader could see, as text (dates excluded)."""
    for ws in wb.worksheets:
        yield ws.title
        yield from (text for _, _, text in _header_footer_texts(ws))
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if value is not None and not isinstance(value, date):
                    yield str(value)


def find_leaks(original: Workbook, redacted: Workbook, names: list[str]) -> list[str]:
    return text_leaks(workbook_texts(original), workbook_texts(redacted), names)


def redact_xlsx_file(
    source: Path, target: Path, names: list[str], rng: random.Random | None = None
) -> None:
    """Write a redacted copy of `source` to `target`, or raise `LeakError` and write nothing."""
    check_paths(source, target)
    rng = rng or random.Random()
    original = load_workbook(source, data_only=True)
    leaks: list[str] = []
    for attempt in range(1, MAX_ATTEMPTS + 1):
        redacted = load_workbook(source, data_only=True)
        redact_workbook(redacted, Redactor(names, rng))
        leaks = find_leaks(original, redacted, names)
        if not leaks:
            target.parent.mkdir(parents=True, exist_ok=True)
            redacted.save(target)
            logger.info("Redacted copy saved", extra={"target": str(target), "attempt": attempt})
            return
    raise leak_error(leaks)
