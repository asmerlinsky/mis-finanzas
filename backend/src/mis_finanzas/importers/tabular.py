"""Generic CSV/Excel importer driven by a saved column mapping.

Works for any export that is a table with a header row, so Santander and Mercado Pago
movement exports import before a dedicated parser exists. Old binary `.xls` files are not
supported; re-save them as `.xlsx` or `.csv`.
"""

import csv
import io
from datetime import date, datetime
from decimal import Decimal
from pathlib import PurePath
from typing import Self

from openpyxl import load_workbook
from pydantic import BaseModel, Field, model_validator

from mis_finanzas.core.money import parse_amount
from mis_finanzas.importers.base import ImportFileError, ParsedRow, ParseResult, RowError

Cell = str | int | float | datetime | date | None
TableRow = tuple[int, dict[str, Cell]]

EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
CSV_SUFFIXES = {".csv", ".txt", ".tsv"}


class TabularMapping(BaseModel):
    """How to read one source's export. Column names are matched after trimming spaces."""

    header_row: int = Field(default=1, ge=1, description="1-based line of the header")
    delimiter: str | None = Field(default=None, description="CSV only; None = detect")
    date_column: str
    date_format: str = "%d/%m/%Y"
    description_columns: list[str] = Field(min_length=1)
    amount_column: str | None = None
    debit_column: str | None = None
    credit_column: str | None = None
    decimal_separator: str = ","
    thousands_separator: str = "."
    negate: bool = Field(default=False, description="Flip signs, e.g. card charges listed >0")

    @model_validator(mode="after")
    def _one_amount_mode(self) -> Self:
        has_split = self.debit_column is not None or self.credit_column is not None
        if (self.amount_column is None) == (not has_split):
            raise ValueError("set either amount_column or debit_column/credit_column")
        return self

    def required_columns(self) -> list[str]:
        optional = [self.amount_column, self.debit_column, self.credit_column]
        return [self.date_column, *self.description_columns, *(c for c in optional if c)]


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ImportFileError("file is not UTF-8 or Windows-1252 text")


def _csv_lines(content: bytes, delimiter: str | None) -> list[list[str]]:
    text = _decode(content)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|").delimiter
        except csv.Error:
            delimiter = ";"
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _excel_cell(value: object) -> Cell:
    """Keep the types the converters understand; anything else becomes text."""
    if value is None or isinstance(value, str | int | float | datetime | date):
        return value
    return str(value)


def excel_lines(content: bytes) -> list[list[Cell]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError(f"not a readable Excel file: {exc}") from exc
    sheet = workbook.worksheets[0]
    lines = [[_excel_cell(c) for c in row] for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return lines


def read_table(content: bytes, filename: str, mapping: TabularMapping) -> list[TableRow]:
    """Rows below the header as `(line_number, {column: cell})`."""
    suffix = PurePath(filename).suffix.lower()
    if suffix in EXCEL_SUFFIXES:
        lines = excel_lines(content)
    elif suffix in CSV_SUFFIXES:
        lines = [list[Cell](row) for row in _csv_lines(content, mapping.delimiter)]
    else:
        raise ImportFileError(f"unsupported file type {suffix!r}; use .csv or .xlsx")
    if len(lines) < mapping.header_row:
        raise ImportFileError(f"file has no line {mapping.header_row} for the header")
    header = [str(c).strip() if c is not None else "" for c in lines[mapping.header_row - 1]]
    missing = [c for c in mapping.required_columns() if c not in header]
    if missing:
        raise ImportFileError(f"columns not found in header: {', '.join(missing)}")
    start = mapping.header_row + 1
    return [
        (start + i, dict(zip(header, line, strict=False)))
        for i, line in enumerate(lines[mapping.header_row :])
    ]


def _is_blank(cell: Cell) -> bool:
    return cell is None or (isinstance(cell, str) and cell.strip() in ("", "-"))


def _to_date(cell: Cell, date_format: str) -> date:
    if isinstance(cell, datetime):
        return cell.date()
    if isinstance(cell, date):
        return cell
    if isinstance(cell, str):
        return datetime.strptime(cell.strip(), date_format).date()
    raise ValueError(f"not a date: {cell!r}")


def _to_amount(cell: Cell, mapping: TabularMapping) -> Decimal:
    if _is_blank(cell):
        return Decimal(0)
    if isinstance(cell, int | float) and not isinstance(cell, bool):
        return Decimal(str(cell))
    if isinstance(cell, str):
        return parse_amount(cell, mapping.decimal_separator, mapping.thousands_separator)
    raise ValueError(f"not an amount: {cell!r}")


def _row_amount(values: dict[str, Cell], mapping: TabularMapping) -> Decimal:
    if mapping.amount_column is not None:
        amount = _to_amount(values.get(mapping.amount_column), mapping)
    else:
        credit = _to_amount(values.get(mapping.credit_column or ""), mapping)
        debit = _to_amount(values.get(mapping.debit_column or ""), mapping)
        amount = abs(credit) - abs(debit)
    return -amount if mapping.negate else amount


def _is_empty_row(values: dict[str, Cell], mapping: TabularMapping) -> bool:
    amount_columns = [mapping.amount_column, mapping.debit_column, mapping.credit_column]
    keys = [mapping.date_column, *(c for c in amount_columns if c)]
    return all(_is_blank(values.get(k)) for k in keys)


def convert_row(line_number: int, values: dict[str, Cell], mapping: TabularMapping) -> ParsedRow:
    """Raises `ValueError` with a readable message when a cell can't be converted."""
    try:
        row_date = _to_date(values.get(mapping.date_column), mapping.date_format)
    except ValueError as exc:
        raise ValueError(f"bad date in {mapping.date_column!r}: {exc}") from exc
    amount = _row_amount(values, mapping)
    parts = [str(values.get(c) or "").strip() for c in mapping.description_columns]
    description = " ".join(p for p in parts if p)
    if not description:
        raise ValueError("empty description")
    return ParsedRow(line_number=line_number, date=row_date, amount=amount, description=description)


class GenericTabularImporter:
    name = "generic_tabular"
    version = "1"

    def __init__(self, mapping: TabularMapping) -> None:
        self.mapping = mapping

    def parse(self, content: bytes, filename: str) -> ParseResult:
        result = ParseResult()
        for line_number, values in read_table(content, filename, self.mapping):
            if _is_empty_row(values, self.mapping):
                continue
            try:
                result.rows.append(convert_row(line_number, values, self.mapping))
            except ValueError as exc:
                result.errors.append(RowError(line_number=line_number, message=str(exc)))
        return result
