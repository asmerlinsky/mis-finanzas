"""Redaction of synthetic statements: formats survive, identifiers and names don't."""

import random
import re
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from mis_finanzas.tools.redact import main
from mis_finanzas.tools.redact_xlsx import find_leaks, redact_workbook, redact_xlsx_file
from mis_finanzas.tools.redaction import LeakError, Redactor

NAME = "Juana Pérez"
CBU = "0720123456789012345678"


def _statement() -> Workbook:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Cuenta 4567890"
    header = ws.oddHeader
    assert header is not None
    header.center.text = f"Titular {NAME}"
    ws.append([f"Titular: {NAME}", f"CBU {CBU}"])
    ws.append(["Fecha", "Concepto", "Importe", "Fecha texto"])
    ws.append([datetime(2026, 3, 2), "TRANSF 20304050607", -12345.67, "02/03/2026"])
    ws.append([datetime(2026, 3, 3), "Pago tarjeta", 1500, "2026-03-03 10:15"])
    ws.merge_cells("E1:F1")
    ws["A3"].comment = Comment(f"nota de {NAME}", "autor")
    hidden = wb.create_sheet("Oculta")
    hidden.sheet_state = "hidden"
    hidden.append([f"Cuenta 99887766 de {NAME}"])
    wb.properties.creator = NAME
    return wb


def _save(wb: Workbook, path: Path) -> Path:
    wb.save(path)
    return path


def _digit_shape(text: str) -> str:
    return re.sub(r"\d", "9", text)


def test_redact_keeps_format_dates_and_types() -> None:
    wb = _statement()
    redact_workbook(wb, Redactor([NAME], random.Random(1)))
    ws = wb.worksheets[0]
    assert ws["A3"].value == datetime(2026, 3, 2)
    assert ws["D3"].value == "02/03/2026"
    assert ws["D4"].value == "2026-03-03 10:15"
    assert isinstance(ws["C3"].value, float) and isinstance(ws["C4"].value, int)
    assert _digit_shape(str(ws["C3"].value)) == _digit_shape("-12345.67")
    assert ws["A1"].value == "Titular: XXXX"
    assert ws["A3"].comment is None
    assert wb.properties.creator != NAME


def test_redact_reaches_titles_headers_and_hidden_sheets() -> None:
    original, redacted = _statement(), _statement()
    redact_workbook(redacted, Redactor([NAME], random.Random(2)))
    assert find_leaks(original, redacted, [NAME]) == []
    header = redacted.worksheets[0].oddHeader
    assert header is not None and "XXXX" in str(header.center.text)
    assert redacted.worksheets[1].sheet_state == "hidden"


def test_find_leaks_reports_surviving_identifiers_and_name_variants() -> None:
    original, redacted = _statement(), _statement()
    redact_workbook(redacted, Redactor([], random.Random(3)))
    redacted.worksheets[0]["B5"] = f"copia {CBU}"
    leaks = find_leaks(original, redacted, ["juana perez"])
    assert CBU in leaks and "juana perez" in leaks


def test_redact_file_writes_clean_copy_and_keeps_original(tmp_path: Path) -> None:
    source = _save(_statement(), tmp_path / "mov.xlsx")
    before = source.read_bytes()
    target = tmp_path / "samples" / "mov.redacted.xlsx"
    redact_xlsx_file(source, target, [NAME], random.Random(4))
    assert source.read_bytes() == before
    assert find_leaks(load_workbook(source), load_workbook(target), [NAME]) == []


def test_name_parts_are_masked_in_any_order_and_accent() -> None:
    redactor = Redactor(["Juana Pérez"], random.Random(0))
    assert redactor.text("PEREZ, JUANA") == "XXXX, XXXX"
    assert redactor.text("Transf. a Perez") == "Transf. a XXXX"
    assert redactor.text("Juanita") == "Juanita"


def test_redact_file_writes_nothing_when_a_name_survives(tmp_path: Path) -> None:
    wb = _statement()
    wb.worksheets[0]["B6"] = "SraJuanaPérez"  # run together: no single word to mask
    source = _save(wb, tmp_path / "mov.xlsx")
    target = tmp_path / "out.xlsx"
    with pytest.raises(LeakError):
        redact_xlsx_file(source, target, ["Juana Perez"], random.Random(5))
    assert not target.exists()


def test_redact_file_refuses_to_overwrite_the_source(tmp_path: Path) -> None:
    source = _save(_statement(), tmp_path / "mov.xlsx")
    with pytest.raises(ValueError):
        redact_xlsx_file(source, source, [NAME])


def test_main_defaults_to_samples_and_rejects_other_types(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    source = _save(_statement(), tmp_path / "mov.xlsx")
    assert main([str(source), "--name", NAME]) == 0
    assert (tmp_path / "samples" / "mov.redacted.xlsx").exists()
    assert main([str(tmp_path / "mov.ods")]) == 2
    assert main([str(tmp_path / "missing.xlsx")]) == 1


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("A agust&#237;n sanc / varios", "A XXXX XXXX / varios"),
        ("De Juana Pere merlins", "De XXXX XXXX merlins"),
        ("Per", "Per"),  # 3 letters: too short to count as a truncation
    ],
)
def test_truncated_names_and_entities_are_masked(text: str, expected: str) -> None:
    redactor = Redactor(["Agustín Sánchez", "Juana Pérez"], random.Random(0))
    assert redactor.text(text) == expected


def test_leak_check_catches_truncations_and_entities() -> None:
    from mis_finanzas.tools.redaction import text_leaks

    leaks = text_leaks([], ["Transf a juana pere", "agust&#237;n"], ["Juana Pérez", "Agustín"])
    assert "JUANA" in leaks and "PERE" in leaks and "AGUSTIN" in leaks
